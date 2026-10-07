#!/usr/bin/env python3
"""Parsea el Excel de fuentes autorizadas a CSVs de semillas para el pipeline.

El Excel ("URLs autorizadas para web scraping OFICIAL.xlsx") tiene formato
libre: filas con [Institución | texto-con-URLs], filas de continuación con
solo URLs, URLs con etiqueta ("Plataforma X: https://..."), URLs desnudas
(www.sitio.cl, demre.cl) y varias URLs por celda.

Salida (formato que `main.py --url-file` consume; la Institución viaja como
extra.seed_context en cada documento recolectado):

    seeds_autorizadas_homes.csv   portales institucionales (ruta / o vacía)
                                  -> crawlear en modo secciones
    seeds_autorizadas_deep.csv    URLs profundas (actas, noticias, listados)
                                  -> crawlear en modo clásico

Uso:
    python scripts/parse_authorized_sources.py <excel> [--out-dir configs/seeds]
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

URL_PAT = re.compile(r"https?://[^\s,;\"'”]+", re.I)
BARE_DOMAIN_PAT = re.compile(
    r"(?<![\w/.-])((?:[a-z0-9-]+\.)+(?:cl|pe|bo|ar|co|cr|cu|ec|sv|gt|hn|mx|ni|"
    r"pa|py|do|uy|ve|br|int|org|com|net|edu|gob|gov|travel))(?![\w.-])", re.I)


def clean_url(u: str) -> str | None:
    u = u.strip().rstrip(".,;:)»”\"'")
    if not u.startswith("http"):
        u = "https://" + u
    p = urlparse(u)
    if not p.netloc or "." not in p.netloc:
        return None
    return u


def extract_urls(text: str) -> list[str]:
    urls = [m.group(0) for m in URL_PAT.finditer(text)]
    # dominios desnudos fuera de las URLs ya encontradas
    remainder = URL_PAT.sub(" ", text)
    urls += [m.group(1) for m in BARE_DOMAIN_PAT.finditer(remainder)]
    out = []
    for u in urls:
        cu = clean_url(u)
        if cu:
            out.append(cu)
    return out


def is_home(url: str) -> bool:
    p = urlparse(url)
    return p.path.strip("/") == "" and not p.query


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("excel", type=Path)
    ap.add_argument("--out-dir", type=Path, default=Path("configs/seeds"))
    args = ap.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(args.excel, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows_out: list[tuple[str, str]] = []      # (institucion, url)
    seen: set[str] = set()
    institucion = ""
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if not cells:
            continue
        # celda(s) sin URL al inicio = nombre de institución (nueva o encabezado)
        joined = " ".join(cells)
        first = cells[0]
        if not URL_PAT.search(first) and not BARE_DOMAIN_PAT.search(first):
            if first.lower().startswith("institución"):
                continue                      # fila de encabezado
            institucion = first
        for url in extract_urls(joined):
            key = url.rstrip("/").lower()
            if key in seen:
                continue
            seen.add(key)
            rows_out.append((institucion, url))

    homes = [(i, u) for i, u in rows_out if is_home(u)]
    deep = [(i, u) for i, u in rows_out if not is_home(u)]

    for name, subset in (("seeds_autorizadas_homes.csv", homes),
                         ("seeds_autorizadas_deep.csv", deep)):
        path = args.out_dir / name
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["Institucion", "URL"])
            w.writerows(subset)
        print(f"{path}: {len(subset)} semillas")

    doms = {urlparse(u).netloc for _, u in rows_out}
    print(f"Total: {len(rows_out)} URLs únicas en {len(doms)} dominios, "
          f"{len({i for i, _ in rows_out if i})} instituciones")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
